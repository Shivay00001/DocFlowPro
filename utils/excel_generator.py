"""
Excel Generator for DocFlow Pro
Generates Excel exports for documents and invoices
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from typing import List, Dict
import json


class ExcelGenerator:
    """Generate Excel spreadsheets for data export"""
    
    def __init__(self):
        self.header_fill = PatternFill(start_color="1a237e", end_color="1a237e", fill_type="solid")
        self.header_font = Font(bold=True, color="FFFFFF", size=12)
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def _style_header_row(self, ws, row_num: int, col_count: int):
        """Apply styling to header row"""
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self.border
    
    def _auto_adjust_columns(self, ws):
        """Auto-adjust column widths"""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def generate_documents_excel(self, documents: List[Dict], filename: str) -> bool:
        """Generate Excel export of documents"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Documents"
            
            # Add metadata
            ws['A1'] = "Documents Export"
            ws['A1'].font = Font(bold=True, size=14)
            ws['A2'] = f"Exported on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            ws['A3'] = f"Total Documents: {len(documents)}"
            
            # Headers
            headers = ['ID', 'Title', 'Description', 'Category', 'Tags', 'File Type', 
                      'File Size (bytes)', 'Created Date', 'Updated Date']
            
            for col, header in enumerate(headers, 1):
                ws.cell(row=5, column=col, value=header)
            
            self._style_header_row(ws, 5, len(headers))
            
            # Data rows
            for row_idx, doc in enumerate(documents, 6):
                ws.cell(row=row_idx, column=1, value=doc.get('id'))
                ws.cell(row=row_idx, column=2, value=doc.get('title', ''))
                ws.cell(row=row_idx, column=3, value=doc.get('description', ''))
                ws.cell(row=row_idx, column=4, value=doc.get('category', ''))
                ws.cell(row=row_idx, column=5, value=doc.get('tags', ''))
                ws.cell(row=row_idx, column=6, value=doc.get('file_type', ''))
                ws.cell(row=row_idx, column=7, value=doc.get('file_size', 0))
                ws.cell(row=row_idx, column=8, value=doc.get('created_at', ''))
                ws.cell(row=row_idx, column=9, value=doc.get('updated_at', ''))
                
                # Apply borders
                for col in range(1, len(headers) + 1):
                    ws.cell(row=row_idx, column=col).border = self.border
            
            # Auto-adjust columns
            self._auto_adjust_columns(ws)
            
            # Save workbook
            wb.save(filename)
            return True
            
        except Exception as e:
            print(f"Error generating documents Excel: {e}")
            return False
    
    def generate_invoices_excel(self, invoices: List[Dict], filename: str) -> bool:
        """Generate Excel export of invoices"""
        try:
            wb = Workbook()
            
            # Summary sheet
            ws_summary = wb.active
            ws_summary.title = "Summary"
            
            ws_summary['A1'] = "Invoices Export Summary"
            ws_summary['A1'].font = Font(bold=True, size=14)
            ws_summary['A2'] = f"Exported on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            ws_summary['A3'] = f"Total Invoices: {len(invoices)}"
            
            total_amount = sum(inv.get('total_amount', 0) for inv in invoices)
            ws_summary['A4'] = f"Total Amount: ₹{total_amount:.2f}"
            
            # Invoices sheet
            ws = wb.create_sheet("Invoices")
            
            # Headers
            headers = ['ID', 'Invoice Number', 'Client Name', 'Client Email', 'Invoice Date', 
                      'Due Date', 'Subtotal', 'Tax Rate (%)', 'Tax Amount', 'Total Amount', 
                      'Status', 'Notes', 'Created Date']
            
            for col, header in enumerate(headers, 1):
                ws.cell(row=1, column=col, value=header)
            
            self._style_header_row(ws, 1, len(headers))
            
            # Data rows
            for row_idx, inv in enumerate(invoices, 2):
                ws.cell(row=row_idx, column=1, value=inv.get('id'))
                ws.cell(row=row_idx, column=2, value=inv.get('invoice_number', ''))
                ws.cell(row=row_idx, column=3, value=inv.get('client_name', ''))
                ws.cell(row=row_idx, column=4, value=inv.get('client_email', ''))
                ws.cell(row=row_idx, column=5, value=inv.get('invoice_date', ''))
                ws.cell(row=row_idx, column=6, value=inv.get('due_date', ''))
                ws.cell(row=row_idx, column=7, value=inv.get('subtotal', 0))
                ws.cell(row=row_idx, column=8, value=inv.get('tax_rate', 0))
                ws.cell(row=row_idx, column=9, value=inv.get('tax_amount', 0))
                ws.cell(row=row_idx, column=10, value=inv.get('total_amount', 0))
                ws.cell(row=row_idx, column=11, value=inv.get('status', ''))
                ws.cell(row=row_idx, column=12, value=inv.get('notes', ''))
                ws.cell(row=row_idx, column=13, value=inv.get('created_at', ''))
                
                # Apply borders
                for col in range(1, len(headers) + 1):
                    ws.cell(row=row_idx, column=col).border = self.border
            
            # Auto-adjust columns
            self._auto_adjust_columns(ws)
            
            # Items sheet
            ws_items = wb.create_sheet("Invoice Items")
            
            item_headers = ['Invoice ID', 'Invoice Number', 'Item Description', 'Quantity', 'Rate', 'Amount']
            for col, header in enumerate(item_headers, 1):
                ws_items.cell(row=1, column=col, value=header)
            
            self._style_header_row(ws_items, 1, len(item_headers))
            
            item_row = 2
            for inv in invoices:
                items = inv.get('items', [])
                if isinstance(items, str):
                    items = json.loads(items)
                
                for item in items:
                    ws_items.cell(row=item_row, column=1, value=inv.get('id'))
                    ws_items.cell(row=item_row, column=2, value=inv.get('invoice_number', ''))
                    ws_items.cell(row=item_row, column=3, value=item.get('description', ''))
                    ws_items.cell(row=item_row, column=4, value=item.get('quantity', 0))
                    ws_items.cell(row=item_row, column=5, value=item.get('rate', 0))
                    amount = item.get('quantity', 0) * item.get('rate', 0)
                    ws_items.cell(row=item_row, column=6, value=amount)
                    
                    # Apply borders
                    for col in range(1, len(item_headers) + 1):
                        ws_items.cell(row=item_row, column=col).border = self.border
                    
                    item_row += 1
            
            self._auto_adjust_columns(ws_items)
            
            # Save workbook
            wb.save(filename)
            return True
            
        except Exception as e:
            print(f"Error generating invoices Excel: {e}")
            return False
    
    def generate_complete_export_excel(self, export_data: Dict, filename: str) -> bool:
        """Generate complete data export with multiple sheets"""
        try:
            wb = Workbook()
            
            # Overview sheet
            ws_overview = wb.active
            ws_overview.title = "Overview"
            
            ws_overview['A1'] = "DocFlow Pro - Complete Data Export"
            ws_overview['A1'].font = Font(bold=True, size=16)
            ws_overview['A2'] = f"Exported on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            
            stats = export_data.get('statistics', {})
            ws_overview['A4'] = "Statistics:"
            ws_overview['A4'].font = Font(bold=True, size=12)
            ws_overview['A5'] = f"Total Documents: {stats.get('total_documents', 0)}"
            ws_overview['A6'] = f"Total Invoices: {stats.get('total_invoices', 0)}"
            ws_overview['A7'] = f"Total Invoice Amount: ₹{stats.get('total_invoice_amount', 0):.2f}"
            
            # Export documents if available
            documents = export_data.get('documents', [])
            if documents:
                ws_docs = wb.create_sheet("Documents")
                headers = ['ID', 'Title', 'Category', 'Tags', 'Created Date']
                
                for col, header in enumerate(headers, 1):
                    ws_docs.cell(row=1, column=col, value=header)
                self._style_header_row(ws_docs, 1, len(headers))
                
                for row_idx, doc in enumerate(documents, 2):
                    ws_docs.cell(row=row_idx, column=1, value=doc.get('id'))
                    ws_docs.cell(row=row_idx, column=2, value=doc.get('title', ''))
                    ws_docs.cell(row=row_idx, column=3, value=doc.get('category', ''))
                    ws_docs.cell(row=row_idx, column=4, value=doc.get('tags', ''))
                    ws_docs.cell(row=row_idx, column=5, value=doc.get('created_at', ''))
                
                self._auto_adjust_columns(ws_docs)
            
            # Export invoices if available
            invoices = export_data.get('invoices', [])
            if invoices:
                ws_inv = wb.create_sheet("Invoices")
                headers = ['Invoice Number', 'Client', 'Date', 'Total Amount', 'Status']
                
                for col, header in enumerate(headers, 1):
                    ws_inv.cell(row=1, column=col, value=header)
                self._style_header_row(ws_inv, 1, len(headers))
                
                for row_idx, inv in enumerate(invoices, 2):
                    ws_inv.cell(row=row_idx, column=1, value=inv.get('invoice_number', ''))
                    ws_inv.cell(row=row_idx, column=2, value=inv.get('client_name', ''))
                    ws_inv.cell(row=row_idx, column=3, value=inv.get('invoice_date', ''))
                    ws_inv.cell(row=row_idx, column=4, value=inv.get('total_amount', 0))
                    ws_inv.cell(row=row_idx, column=5, value=inv.get('status', ''))
                
                self._auto_adjust_columns(ws_inv)
            
            # Save workbook
            wb.save(filename)
            return True
            
        except Exception as e:
            print(f"Error generating complete export Excel: {e}")
            return False
